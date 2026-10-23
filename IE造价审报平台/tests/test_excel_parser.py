"""
Excel 解析器单元测试
"""
import os
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

from services.excel_parser import (
    safe_float,
    detect_column_mapping,
    get_unit_from_name,
    parse_excel_file,
    parse_excel_file_multi_sheet,
)

BACKEND_DIR = Path(__file__).resolve().parent.parent / "backend"
sys.path.insert(0, str(BACKEND_DIR))


class TestSafeFloat:
    def test_int_string(self):
        assert safe_float("100") == 100.0

    def test_float_string(self):
        assert safe_float("3.14") == 3.14

    def test_with_chinese_chars(self):
        # 含中文的非纯数字字符串无法转换
        assert safe_float("100元") is None

    def test_none(self):
        assert safe_float(None) is None

    def test_empty_string(self):
        assert safe_float("") is None

    def test_non_numeric(self):
        assert safe_float("abc") is None

    def test_negative(self):
        assert safe_float("-50.5") == -50.5


class TestGetUnitFromName:
    def test_m3(self):
        assert get_unit_from_name("挖一般土方") == "m3"

    def test_m2(self):
        assert get_unit_from_name("路面铺装面积") == "m2"

    def test_m(self):
        assert get_unit_from_name("管道长度") == "m"

    def test_ton(self):
        assert get_unit_from_name("钢筋t") == "t"

    def test_no_match_returns_empty(self):
        # 无单位关键词时返回空字符串
        assert get_unit_from_name("措施项目费") == ""


class TestParseExcelFile:
    def _make_simple_excel(self, path):
        wb = Workbook()
        ws = wb.active
        ws.title = "道路工程"
        # 表头（解析器支持 送审/审定/审核增减-金额 关键字）
        ws.append(["序号", "项目名称", "单位", "送审工程量", "送审单价", "送审合价",
                   "审定工程量", "审定单价", "审定合价", "审核增减-金额"])
        ws.append([1, "挖一般土方", "m3", 5000, 10, 50000, 4800, 10, 48000, -2000])
        ws.append([2, "管理费", "元", 1, 1000, 1000, 1, 800, 800, -200])
        wb.save(path)

    def test_parse_single_sheet(self, tmp_path):
        path = tmp_path / "test.xlsx"
        self._make_simple_excel(str(path))
        result = parse_excel_file(str(path))
        assert result["engineering_type"] == "道路工程"
        assert "items" in result
        items = result["items"]
        assert len(items) == 2
        item0 = items[0]
        assert item0["name"] == "挖一般土方"
        assert item0["submitted_qty"] == 5000
        assert item0["audited_qty"] == 4800
        assert item0["adjustment"] == -2000

    def test_parse_multi_sheet(self, tmp_path):
        # 创建多 sheet 文件
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "道路工程"
        ws1.append(["序号", "项目名称", "单位", "送审工程量", "送审单价", "送审合价",
                    "审定工程量", "审定单价", "审定合价", "审核增减-金额"])
        ws1.append([1, "挖土方", "m3", 100, 5, 500, 90, 5, 450, -50])
        ws2 = wb.create_sheet("排水工程")
        ws2.append(["序号", "项目名称", "单位", "送审工程量", "送审单价", "送审合价",
                    "审定工程量", "审定单价", "审定合价", "审核增减-金额"])
        ws2.append([1, "管道", "m", 200, 100, 20000, 180, 100, 18000, -2000])
        path = tmp_path / "multi.xlsx"
        wb.save(str(path))

        results = parse_excel_file_multi_sheet(str(path))
        assert len(results) >= 2
        types = {r["engineering_type"] for r in results}
        assert "道路工程" in types
        assert "排水工程" in types

    def test_nonexistent_file_raises(self):
        with pytest.raises((FileNotFoundError, ValueError, Exception)):
            parse_excel_file("/nonexistent/path/to/file.xlsx")


class TestMultiLevelHeader:
    """多级表头解析测试（送审工程/审结工程/审计结果 组表头）"""

    def _make_multilevel_excel(self, path):
        """构造多级表头 Excel：
        R3: 组表头（F=送审工程, J=审结工程, M=审计结果）
        R4: 子表头 金额（元）
        R5: 叶子表头（综合单价/合价/审增金额）
        R6+: 数据行
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "道路工程"
        # Row 3: 基础列 + 组表头
        ws.cell(3, 1, "序号")
        ws.cell(3, 3, "项目名称")
        ws.cell(3, 4, "计量单位")
        ws.cell(3, 6, "送审工程")    # F3 → group 送审 (F-I = 6-9)
        ws.cell(3, 10, "审结工程")   # J3 → group 审结 (J-L = 10-12)
        ws.cell(3, 13, "审计结果")   # M3 → group 审计结果 (M-N = 13-14)
        # Row 4: 子表头
        ws.cell(4, 7, "金额（元）")   # G4
        ws.cell(4, 9, "金额（元）")   # I4
        ws.cell(4, 11, "金额（元）")  # K4
        ws.cell(4, 12, "金额（元）")  # L4
        ws.cell(4, 14, "金额（元）")  # N4
        # Row 5: 叶子表头
        ws.cell(5, 6, "工程量")      # F5 → submitted_qty
        ws.cell(5, 7, "综合单价")    # G5 → submitted_price
        ws.cell(5, 9, "合价")        # I5 → submitted_total
        ws.cell(5, 10, "工程量")     # J5 → audited_qty
        ws.cell(5, 11, "综合单价")   # K5 → audited_price
        ws.cell(5, 12, "合价")       # L5 → audited_total
        ws.cell(5, 14, "审增金额")   # N5 → adjustment（别名）
        # Row 6: 数据行
        ws.cell(6, 1, 1)
        ws.cell(6, 3, "挖一般土方")
        ws.cell(6, 4, "m3")
        ws.cell(6, 6, 100)       # submitted_qty
        ws.cell(6, 7, 10)        # submitted_price
        ws.cell(6, 9, 1000)      # submitted_total
        ws.cell(6, 10, 90)       # audited_qty
        ws.cell(6, 11, 10)       # audited_price
        ws.cell(6, 12, 900)      # audited_total
        ws.cell(6, 14, -100)     # adjustment
        wb.save(path)

    def test_multilevel_header_parsing(self, tmp_path):
        """多级表头下单价/合价/调整金额正确解析（核心修复点）"""
        path = tmp_path / "multilevel.xlsx"
        self._make_multilevel_excel(str(path))
        result = parse_excel_file(str(path))
        assert result["engineering_type"] == "道路工程"
        items = result["items"]
        assert len(items) == 1
        item = items[0]
        assert item["name"] == "挖一般土方"
        assert item["unit"] == "m3"
        # 核心断言：多级表头下藏匿的单价和合价
        assert item["submitted_price"] == 10    # G 列
        assert item["audited_price"] == 10      # K 列
        assert item["submitted_total"] == 1000  # I 列
        assert item["audited_total"] == 900     # L 列
        assert item["adjustment"] == -100       # N 列（审增金额别名）
        # 工程量
        assert item["submitted_qty"] == 100
        assert item["audited_qty"] == 90

    def test_adjustment_alias_recognition(self, tmp_path):
        """叶子里含'审增金额'别名时正确识别为 adjustment 列"""
        wb = Workbook()
        ws = wb.active
        ws.title = "排水工程"
        ws.append(["序号", "项目名称", "单位", "送审合价", "审定合价", "审增金额"])
        ws.append([1, "管道", "m", 20000, 18000, -2000])
        path = tmp_path / "alias.xlsx"
        wb.save(str(path))
        result = parse_excel_file(str(path))
        items = result["items"]
        assert len(items) == 1
        assert items[0]["adjustment"] == -2000

    def test_column_mapping_override(self, tmp_path):
        """手工列映射 override 生效，未 override 的字段仍走自动检测"""
        path = tmp_path / "override.xlsx"
        self._make_multilevel_excel(str(path))
        # 把 name 列强制指向列 1（序号），验证 override 覆盖自动检测
        result = parse_excel_file(
            str(path),
            column_mapping_override={"道路工程": {"name": 1}},
        )
        items = result["items"]
        assert len(items) >= 1
        # name 列被改为列 1（序号），数据行的序号值是 1
        assert str(items[0]["name"]) == "1"
        # 未 override 的字段仍走自动检测
        assert items[0]["submitted_price"] == 10
        assert items[0]["audited_total"] == 900

    def test_adjustment_subgroup_header_excluded(self, tmp_path):
        """审计结果组下"审核增减"作为子组标题（跨 M:N 合并，值只在 M）时，
        adjustment 应选 N 列（金额叶子），而非 M 列（工程量叶子）。
        这是 XX路项目.xlsx 报告的 bug：子组标题"审核增减"被当作 M 列叶子文本拼接，
        命中 _ADJUSTMENT_ALIASES 后误选 M 列。"""
        wb = Workbook()
        ws = wb.active
        ws.title = "道路工程"
        # Row 3: 基础列 + 组表头
        ws.cell(3, 1, "序号")
        ws.cell(3, 3, "项目名称")
        ws.cell(3, 5, "计量单位")
        ws.cell(3, 6, "送审工程")    # F3 → 送审 (F-I)
        ws.cell(3, 10, "审结工程")   # J3 → 审结 (J-L)
        ws.cell(3, 13, "审计结果")   # M3 → 审计结果 (M-N)
        # Row 4: 子组标题"审核增减"在 M（跨 M:N 合并，N 为空）
        ws.cell(4, 6, "工程量")      # F4
        ws.cell(4, 7, "金额（元）")  # G4
        ws.cell(4, 10, "工程量")     # J4
        ws.cell(4, 11, "金额（元）") # K4
        ws.cell(4, 13, "审核增减")   # M4 ← 子组标题（旧代码在此误命中）
        ws.cell(4, 15, "审核说明")   # O4
        # Row 5: 叶子表头
        ws.cell(5, 7, "综合单价")
        ws.cell(5, 9, "合价")
        ws.cell(5, 11, "综合单价")
        ws.cell(5, 12, "合价")
        ws.cell(5, 13, "工程量")     # M5 叶子（工程量差）
        ws.cell(5, 14, "金额")       # N5 叶子（金额差，真正的 adjustment）
        # Row 6: 数据行（复刻 XX路项目 真实数据）
        ws.cell(6, 1, 1)
        ws.cell(6, 3, "挖一般土方")
        ws.cell(6, 5, "m3")
        ws.cell(6, 6, 56583.25)      # submitted_qty
        ws.cell(6, 7, 6.16)          # submitted_price
        ws.cell(6, 9, 348552.82)     # submitted_total
        ws.cell(6, 10, 56757.51)     # audited_qty
        ws.cell(6, 11, 5.94)         # audited_price
        ws.cell(6, 12, 337139.61)    # audited_total
        ws.cell(6, 13, 174.26)       # M6 = 工程量差（不是 adjustment）
        ws.cell(6, 14, -11413.21)    # N6 = 金额差（真正的 adjustment）
        path = tmp_path / "subgroup.xlsx"
        wb.save(str(path))
        result = parse_excel_file(str(path))
        items = result["items"]
        assert len(items) == 1
        item = items[0]
        # 核心断言：adjustment 应取 N 列（-11413.21），而非 M 列（174.26）
        assert item["adjustment"] == -11413.21, (
            f"应选 N 列金额差(-11413.21)，实际得到 {item['adjustment']}（误选了 M 列工程量差）"
        )
        # 工程量差不应被误认为 adjustment
        assert item["submitted_qty"] == 56583.25
        assert item["audited_total"] == 337139.61

    def test_adjustment_prefers_alias_over_generic_amount(self, tmp_path):
        """审计结果组内 M=金额（元）、N=审增金额时，adjustment 应选 N 而非 M。
        这是用户报告的 bug：旧代码先匹配通用"金额"导致 M 抢先。"""
        wb = Workbook()
        ws = wb.active
        ws.title = "道路工程"
        # Row 3: 基础列 + 组表头
        ws.cell(3, 1, "序号")
        ws.cell(3, 3, "项目名称")
        ws.cell(3, 4, "计量单位")
        ws.cell(3, 6, "送审工程")    # F3 → 送审 (F-I)
        ws.cell(3, 10, "审结工程")   # J3 → 审结 (J-L)
        ws.cell(3, 13, "审计结果")   # M3 → 审计结果 (M-N)
        # Row 4: 子表头 —— 关键：M 也有"金额（元）"
        ws.cell(4, 7, "金额（元）")   # G4
        ws.cell(4, 9, "金额（元）")   # I4
        ws.cell(4, 11, "金额（元）")  # K4
        ws.cell(4, 12, "金额（元）")  # L4
        ws.cell(4, 13, "金额（元）")  # M4 ← 旧代码在此命中"金额"，错误选 M
        ws.cell(4, 14, "金额（元）")  # N4
        # Row 5: 叶子表头
        ws.cell(5, 6, "工程量")
        ws.cell(5, 7, "综合单价")
        ws.cell(5, 9, "合价")
        ws.cell(5, 10, "工程量")
        ws.cell(5, 11, "综合单价")
        ws.cell(5, 12, "合价")
        ws.cell(5, 13, "审定金额")   # M5 叶子（审定金额，非 adjustment）
        ws.cell(5, 14, "审增金额")   # N5 叶子（真正的 adjustment）
        # Row 6: 数据行
        ws.cell(6, 1, 1)
        ws.cell(6, 3, "挖一般土方")
        ws.cell(6, 4, "m3")
        ws.cell(6, 6, 100)
        ws.cell(6, 7, 10)
        ws.cell(6, 9, 1000)
        ws.cell(6, 10, 90)
        ws.cell(6, 11, 10)
        ws.cell(6, 12, 900)
        ws.cell(6, 13, 900)     # M6 = 审定金额（不是 adjustment）
        ws.cell(6, 14, -100)    # N6 = 审增金额（真正的 adjustment）
        path = tmp_path / "m_vs_n.xlsx"
        wb.save(str(path))
        result = parse_excel_file(str(path))
        items = result["items"]
        assert len(items) == 1
        item = items[0]
        # 核心断言：adjustment 应取 N 列（-100），而非 M 列（900）
        assert item["adjustment"] == -100, f"应选 N 列审增金额(-100)，实际得到 {item['adjustment']}"
        # 审定金额不应被误认为 adjustment
        assert item["audited_total"] == 900

    def test_adjustment_subgroup_non_keyword(self, tmp_path):
        """子组标题不是"审核增减"而是其他名称（如"差异明细"）时，
        通用结构推断仍能正确排除子组，识别真正的金额叶子列。
        旧代码依赖关键词("审核增减","审增减")会漏判，新代码基于结构推断通用适用。"""
        wb = Workbook()
        ws = wb.active
        ws.title = "道路工程"
        # Row 3: 基础列 + 组表头
        ws.cell(3, 1, "序号")
        ws.cell(3, 3, "项目名称")
        ws.cell(3, 5, "计量单位")
        ws.cell(3, 6, "送审工程")    # F3 → 送审 (F-I)
        ws.cell(3, 10, "审结工程")   # J3 → 审结 (J-L)
        ws.cell(3, 13, "审计结果")   # M3 → 审计结果 (M-N)
        # Row 4: 子组标题"差异明细"在 M（非"审核增减"关键词，跨 M:N 合并，N 为空）
        ws.cell(4, 6, "工程量")
        ws.cell(4, 7, "金额（元）")
        ws.cell(4, 10, "工程量")
        ws.cell(4, 11, "金额（元）")
        ws.cell(4, 13, "差异明细")   # M4 ← 非"审核增减"的子组标题
        # Row 5: 叶子表头
        ws.cell(5, 7, "综合单价")
        ws.cell(5, 9, "合价")
        ws.cell(5, 11, "综合单价")
        ws.cell(5, 12, "合价")
        ws.cell(5, 13, "工程量")     # M5 叶子（工程量差）
        ws.cell(5, 14, "金额")       # N5 叶子（金额差，真正的 adjustment）
        # Row 6: 数据行
        ws.cell(6, 1, 1)
        ws.cell(6, 3, "挖一般土方")
        ws.cell(6, 5, "m3")
        ws.cell(6, 6, 56583.25)
        ws.cell(6, 7, 6.16)
        ws.cell(6, 9, 348552.82)
        ws.cell(6, 10, 56757.51)
        ws.cell(6, 11, 5.94)
        ws.cell(6, 12, 337139.61)
        ws.cell(6, 13, 174.26)       # M6 = 工程量差（不是 adjustment）
        ws.cell(6, 14, -11413.21)    # N6 = 金额差（真正的 adjustment）
        path = tmp_path / "non_keyword_subgroup.xlsx"
        wb.save(str(path))
        result = parse_excel_file(str(path))
        items = result["items"]
        assert len(items) == 1
        item = items[0]
        # 核心断言：adjustment 应取 N 列（-11413.21），而非 M 列（174.26）
        assert item["adjustment"] == -11413.21, (
            f"应选 N 列金额差(-11413.21)，实际得到 {item['adjustment']}（误选了 M 列工程量差）"
        )
        assert item["submitted_qty"] == 56583.25
        assert item["audited_total"] == 337139.61

    def test_single_level_header_adjustment(self, tmp_path):
        """单级表头（仅 1 行表头）下，列名直接为"调整金额"时应正确识别。
        这是测试数据表.xlsx 的场景：R1 是表头行，K 列是"调整金额"。
        旧代码 _scan_group_headers 把"送审数量"/"审定数量"/"调整金额"等叶子
        误判为组表头，导致 adjustment 列识别为 None。"""
        wb = Workbook()
        ws = wb.active
        ws.title = "道路工程"
        # R1: 单级表头（复刻测试数据表.xlsx 结构）
        ws.cell(1, 1, "序号")
        ws.cell(1, 2, "项目名称")
        ws.cell(1, 3, "项目特征")
        ws.cell(1, 4, "单位")
        ws.cell(1, 5, "送审数量")
        ws.cell(1, 6, "送审单价")
        ws.cell(1, 7, "送审总价")
        ws.cell(1, 8, "审定数量")
        ws.cell(1, 9, "审定单价")
        ws.cell(1, 10, "审定总价")
        ws.cell(1, 11, "调整金额")
        # R2: 数据行
        ws.cell(2, 1, 1)
        ws.cell(2, 2, "挖一般土方")
        ws.cell(2, 3, "三类土")
        ws.cell(2, 4, "m3")
        ws.cell(2, 5, 5000)
        ws.cell(2, 6, 25.5)
        ws.cell(2, 7, 127500)
        ws.cell(2, 8, 4800)
        ws.cell(2, 9, 24.8)
        ws.cell(2, 10, 119040)
        ws.cell(2, 11, -8460)
        path = tmp_path / "single_level.xlsx"
        wb.save(str(path))
        result = parse_excel_file(str(path))
        items = result["items"]
        assert len(items) == 1
        item = items[0]
        # 核心断言：单级表头下 adjustment 应正确识别为 K 列（-8460）
        assert item["adjustment"] == -8460, (
            f"应识别 K 列调整金额(-8460)，实际得到 {item['adjustment']}"
        )
        # 其他字段也应正确识别
        assert item["name"] == "挖一般土方"
        assert item["unit"] == "m3"
        assert item["submitted_qty"] == 5000
        assert item["submitted_price"] == 25.5
        assert item["submitted_total"] == 127500
        assert item["audited_qty"] == 4800
        assert item["audited_price"] == 24.8
        assert item["audited_total"] == 119040
