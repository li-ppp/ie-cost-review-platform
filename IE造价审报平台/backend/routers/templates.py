"""
模板管理路由
"""

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Request
from sqlalchemy.orm import Session
from typing import List
import os
import re
import logging
from datetime import datetime

from database import get_db
from models import Template, User, Report
from schemas import TemplateResponse
from config import UPLOADS_DIR
from auth import get_current_user, require_admin
from routers.logs import log_operation

router = APIRouter()
logger = logging.getLogger(__name__)

UPLOAD_DIR = os.path.join(UPLOADS_DIR, "templates")
os.makedirs(UPLOAD_DIR, exist_ok=True)

MAX_LIMIT = 500
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB


def _sanitize_filename(filename: str) -> str:
    """清理文件名，保留中文字符，仅移除不安全的路径字符"""
    safe = re.sub(r'[\\/:*?"<>|\r\n\t]', '_', filename)
    if len(safe) > 100:
        name, ext = os.path.splitext(safe)
        safe = name[:100 - len(ext)] + ext
    return safe


@router.post("/", response_model=TemplateResponse)
async def create_template(
    name: str = Form(...),
    template_type: str = Form(...),
    file: UploadFile = File(...),
    request: Request = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """上传模板"""
    if not file.filename:
        raise HTTPException(status_code=400, detail="未提供文件名")

    file_ext = os.path.splitext(file.filename)[1].lower()
    if file_ext not in ['.docx']:
        raise HTTPException(status_code=400, detail="仅支持 .docx 格式")

    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=400, detail="文件大小不能超过10MB")

    # 验证文件为有效的 Office Open XML（ZIP）格式
    if not content[:4].startswith(b'PK'):
        raise HTTPException(status_code=400, detail="文件格式无效，不是有效的 Word 文档")

    safe_filename = _sanitize_filename(file.filename)
    file_path = os.path.join(UPLOAD_DIR, f"{datetime.now().timestamp()}_{safe_filename}")

    try:
        with open(file_path, "wb") as buffer:
            buffer.write(content)
    except OSError as e:
        logger.error(f"模板文件写入失败: {e}")
        raise HTTPException(status_code=500, detail="文件保存失败")
    finally:
        await file.close()

    template = Template(
        name=name,
        template_type=template_type,
        file_path=file_path,
        variables={}
    )
    try:
        db.add(template)
        db.commit()
        db.refresh(template)
    except Exception as e:
        # DB 提交失败时清理物理文件，避免孤儿文件
        if os.path.exists(file_path):
            try:
                os.remove(file_path)
            except OSError:
                pass
        logger.error(f"模板数据库保存失败: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"模板保存失败: {str(e)}")

    logger.info(f"模板上传成功: {name}, 用户: {getattr(current_user, 'username', 'unknown')}")
    log_operation(db, current_user.id, current_user.username, "上传模板",
                target_type="template", target_id=template.id,
                detail=f"上传模板: {name}, 类型: {template_type}", ip_address=request.client.host)
    return template


@router.get("/", response_model=List[TemplateResponse])
def get_templates(
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """获取模板列表"""
    if skip < 0:
        skip = 0
    limit = max(1, min(limit, MAX_LIMIT))
    templates = db.query(Template).offset(skip).limit(limit).all()
    return templates


@router.get("/{template_id}", response_model=TemplateResponse)
def get_template(
    template_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """获取模板详情"""
    template = db.query(Template).filter(Template.id == template_id).first()
    if not template:
        raise HTTPException(status_code=404, detail="模板不存在")
    return template


@router.delete("/{template_id}")
def delete_template(
    template_id: int,
    request: Request,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """删除模板"""
    template = db.query(Template).filter(Template.id == template_id).first()
    if not template:
        raise HTTPException(status_code=404, detail="模板不存在")

    # 检查是否被报告引用，避免删除后产生孤儿报告
    ref_count = db.query(Report).filter(Report.template_id == template_id).count()
    if ref_count > 0:
        raise HTTPException(
            status_code=400,
            detail=f"模板已被 {ref_count} 个报告引用，无法删除。请先删除相关报告或更换其模板。"
        )

    # 删除物理文件
    if template.file_path and os.path.exists(template.file_path):
        try:
            os.remove(template.file_path)
        except OSError as e:
            logger.warning(f"删除模板文件失败 {template.file_path}: {e}")

    db.delete(template)
    db.commit()
    logger.info(f"模板已删除: {template_id}, 用户: {getattr(current_user, 'username', 'unknown')}")
    log_operation(db, current_user.id, current_user.username, "删除模板",
                target_type="template", target_id=template_id,
                detail=f"删除模板: {template.name}", ip_address=request.client.host)
    return {"message": "模板已删除"}


@router.get("/{template_id}/preview")
def preview_template(
    template_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """预览模板内容"""
    template = db.query(Template).filter(Template.id == template_id).first()
    if not template:
        raise HTTPException(status_code=404, detail="模板不存在")

    if not template.file_path or not os.path.exists(template.file_path):
        raise HTTPException(status_code=404, detail="模板文件不存在")

    try:
        from docx import Document
        doc = Document(template.file_path)

        content = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                content.append(text)

        return {
            "name": template.name,
            "content": "\n".join(content)
        }
    except Exception as e:
        logger.error(f"模板预览失败: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail="模板预览失败")


@router.get("/{template_id}/variables")
def get_template_variables(
    template_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """获取模板变量配置（自动扫描占位符 + 已保存配置）"""
    import re
    from docx import Document

    template = db.query(Template).filter(Template.id == template_id).first()
    if not template:
        raise HTTPException(status_code=404, detail="模板不存在")

    if not template.file_path or not os.path.exists(template.file_path):
        raise HTTPException(status_code=404, detail="模板文件不存在")

    try:
        doc = Document(template.file_path)
        # 扫描段落和表格中的 {{xxx}} 占位符
        placeholders = set()
        pattern = re.compile(r"\{\{([^}]+)\}\}")

        for para in doc.paragraphs:
            for m in pattern.finditer(para.text):
                placeholders.add(m.group(1).strip())

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for m in pattern.finditer(para.text):
                            placeholders.add(m.group(1).strip())

        # 已保存的配置
        saved_vars = template.variables or {}

        # 合并：扫描到的 + 已配置的
        result = []
        for ph in sorted(placeholders):
            cfg = saved_vars.get(ph, {}) if isinstance(saved_vars, dict) else {}
            result.append({
                "placeholder": f"{{{{{ph}}}}}",
                "field_name": cfg.get("field_name", ph),
                "description": cfg.get("description", ""),
                "default_value": cfg.get("default_value", ""),
            })

        # 已配置但模板中已删除的
        if isinstance(saved_vars, dict):
            for k, v in saved_vars.items():
                if k not in placeholders:
                    result.append({
                        "placeholder": f"{{{{{k}}}}}",
                        "field_name": v.get("field_name", k) if isinstance(v, dict) else k,
                        "description": v.get("description", "") if isinstance(v, dict) else "",
                        "default_value": v.get("default_value", "") if isinstance(v, dict) else "",
                    })

        return {"template_id": template_id, "variables": result}
    except Exception as e:
        logger.error(f"获取模板变量失败: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail="获取模板变量失败")


@router.put("/{template_id}/variables")
def update_template_variables(
    template_id: int,
    data: dict,
    request: Request,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin)
):
    """更新模板变量映射配置"""
    template = db.query(Template).filter(Template.id == template_id).first()
    if not template:
        raise HTTPException(status_code=404, detail="模板不存在")

    variables_list = data.get("variables", []) if isinstance(data, dict) else []
    variables_dict = {}
    for item in variables_list:
        ph = item.get("placeholder", "")
        # 去除 {{ }} 包装得到 key
        key = ph.replace("{{", "").replace("}}", "").strip()
        if key:
            variables_dict[key] = {
                "field_name": item.get("field_name", key),
                "description": item.get("description", ""),
                "default_value": item.get("default_value", ""),
            }

    template.variables = variables_dict
    db.commit()
    db.refresh(template)

    logger.info(f"模板变量配置更新: {template_id}, 用户: {getattr(current_user, 'username', 'unknown')}")
    log_operation(db, current_user.id, current_user.username, "更新模板变量",
                target_type="template", target_id=template_id,
                detail=f"配置变量数量: {len(variables_dict)}", ip_address=request.client.host)
    return {"message": "变量配置已保存", "variables_count": len(variables_dict)}


@router.get("/{template_id}/check")
def check_template_compatibility(
    template_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """检查模板结构兼容性，返回兼容性报告"""
    from services.report_generator import validate_template

    template = db.query(Template).filter(Template.id == template_id).first()
    if not template:
        raise HTTPException(status_code=404, detail="模板不存在")

    if not template.file_path or not os.path.exists(template.file_path):
        raise HTTPException(status_code=404, detail="模板文件不存在")

    try:
        result = validate_template(template.file_path)
        logger.info(f"模板兼容性检查: template_id={template_id}, is_compatible={result.get('is_compatible')}")
        return result
    except FileNotFoundError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        logger.error(f"模板检查失败: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"模板检查失败: {str(e)}")
